from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.parsers.base import ParsedSection
from app.parsers.factory import ParserConfigError


@dataclass(frozen=True)
class DocxParagraph:
    text: str
    style_id: str
    style_name: str


@dataclass(frozen=True)
class DocxHeading:
    index: int
    level: int
    title: str
    section_no: str | None


@dataclass(frozen=True)
class NumberingLevel:
    start: int
    num_format: str
    level_text: str


@dataclass(frozen=True)
class DocxBlock:
    kind: str
    paragraph: DocxParagraph | None = None
    table_markdown: str = ""
    paragraph_index: int | None = None


def parse_python_docx_sections(file_path: str, file_name: str) -> list[ParsedSection]:
    suffix = Path(file_name).suffix.lower()
    if suffix != ".docx":
        raise ParserConfigError(
            f"python_docx section parse mode supports only .docx files; got {suffix or 'unknown'}."
        )

    try:
        from docx import Document
    except ImportError as exc:
        raise ParserConfigError("python-docx is not installed.") from exc

    document = Document(file_path)
    paragraph_resolver = NumberingResolver(document)
    paragraphs = [_paragraph_info(paragraph, paragraph_resolver) for paragraph in document.paragraphs]
    paragraphs = [paragraph for paragraph in paragraphs if paragraph.text]
    blocks = _iter_document_blocks(document)
    if not paragraphs and not blocks:
        return []

    body_start_index = _find_body_start_index(paragraphs)
    use_native_headings = any(
        _heading_level_from_style(paragraph.style_id, paragraph.style_name) in {1, 2, 3}
        for paragraph in paragraphs
        if not _is_toc_style(paragraph.style_id, paragraph.style_name)
    )

    roots: list[ParsedSection] = []
    stack: list[ParsedSection] = []
    body_lines_before_first_heading: list[str] = []
    pending_chapter_heading: DocxHeading | None = None

    body_started = body_start_index == 0
    for block in blocks:
        if block.kind == "table":
            if not body_started or not block.table_markdown:
                continue
            if stack:
                _append_content(stack[-1], block.table_markdown)
            else:
                body_lines_before_first_heading.append(block.table_markdown)
            continue

        if block.paragraph is None or block.paragraph_index is None:
            continue
        paragraph = block.paragraph
        index = block.paragraph_index
        text = paragraph.text
        if index < body_start_index:
            continue
        body_started = True
        if _is_toc_style(paragraph.style_id, paragraph.style_name) or _is_toc_label(text):
            continue
        if _looks_like_toc_entry(text):
            continue

        heading = _detect_heading(paragraph, index, use_native_headings)
        if heading:
            if _is_standalone_chapter_no(heading.title):
                pending_chapter_heading = heading
                continue
            if pending_chapter_heading and heading.level <= 2:
                heading = _merge_pending_chapter_heading(pending_chapter_heading, heading)
                pending_chapter_heading = None
            section = ParsedSection(
                level=heading.level,
                title=heading.title,
                section_no=heading.section_no,
            )
            while stack and stack[-1].level >= heading.level:
                stack.pop()
            if stack:
                stack[-1].children.append(section)
            else:
                roots.append(section)
            stack.append(section)
            continue

        if pending_chapter_heading:
            section = ParsedSection(
                level=pending_chapter_heading.level,
                title=pending_chapter_heading.title,
                section_no=pending_chapter_heading.section_no,
            )
            while stack and stack[-1].level >= section.level:
                stack.pop()
            if stack:
                stack[-1].children.append(section)
            else:
                roots.append(section)
            stack.append(section)
            pending_chapter_heading = None

        if stack:
            _append_content(stack[-1], text)
        else:
            body_lines_before_first_heading.append(text)

    if roots and body_lines_before_first_heading:
        prefix = "\n".join(body_lines_before_first_heading).strip()
        first = roots[0]
        first.content = f"{prefix}\n{first.content}".strip() if first.content else prefix
    return roots


def _iter_document_blocks(document) -> list[DocxBlock]:
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[DocxBlock] = []
    paragraph_index = 0
    numbering_resolver = NumberingResolver(document)
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = _paragraph_info(Paragraph(child, document), numbering_resolver)
            if paragraph.text:
                blocks.append(DocxBlock(kind="paragraph", paragraph=paragraph, paragraph_index=paragraph_index))
                paragraph_index += 1
            continue
        if isinstance(child, CT_Tbl):
            table_markdown = _table_to_markdown(Table(child, document))
            if table_markdown:
                blocks.append(DocxBlock(kind="table", table_markdown=table_markdown))
    return blocks


def _paragraph_info(paragraph, numbering_resolver: "NumberingResolver | None" = None) -> DocxParagraph:
    style = paragraph.style
    style_id = getattr(style, "style_id", "") if style else ""
    style_name = getattr(style, "name", "") if style else ""
    text = paragraph.text.strip()
    numbering_label = numbering_resolver.next_label(paragraph) if numbering_resolver else ""
    if numbering_label and text and not text.startswith(numbering_label):
        text = f"{numbering_label} {text}".strip()
    return DocxParagraph(
        text=text,
        style_id=style_id or "",
        style_name=style_name or "",
    )


def _table_to_markdown(table) -> str:
    rows = [
        [_cell_text(cell) for cell in row.cells]
        for row in table.rows
    ]
    rows = _trim_empty_rows(rows)
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = [_escape_table_cell(cell) for cell in normalized[0]]
    separator = ["---"] * width
    body = [[_escape_table_cell(cell) for cell in row] for row in normalized[1:]]
    lines = [
        f"| {' | '.join(header)} |",
        f"| {' | '.join(separator)} |",
    ]
    lines.extend(f"| {' | '.join(row)} |" for row in body)
    return "\n".join(lines)


def _cell_text(cell) -> str:
    return "\n".join(paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip())


def _trim_empty_rows(rows: list[list[str]]) -> list[list[str]]:
    return [row for row in rows if any(cell.strip() for cell in row)]


def _escape_table_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>").strip()


def _append_content(section: ParsedSection, content: str) -> None:
    section.content = f"{section.content}\n\n{content}".strip() if section.content else content


class NumberingResolver:
    def __init__(self, document) -> None:
        self._num_to_abstract: dict[str, str] = {}
        self._levels: dict[tuple[str, str], NumberingLevel] = {}
        self._counters: dict[tuple[str, str], int] = {}
        self._load(document)

    def next_label(self, paragraph) -> str:
        num_id, ilvl = _paragraph_numbering(paragraph)
        if num_id is None or ilvl is None:
            return ""

        abstract_id = self._num_to_abstract.get(num_id)
        if abstract_id is None:
            return ""

        level = self._levels.get((abstract_id, ilvl))
        if level is None or not level.level_text:
            return ""

        current_key = (num_id, ilvl)
        current = self._counters.get(current_key, level.start - 1) + 1
        self._counters[current_key] = current

        current_level = int(ilvl)
        for key in list(self._counters):
            key_num_id, key_ilvl = key
            if key_num_id == num_id and int(key_ilvl) > current_level:
                self._counters.pop(key, None)

        return self._format_level_text(num_id, abstract_id, ilvl, level.level_text).strip()

    def _load(self, document) -> None:
        try:
            numbering = document.part.numbering_part.element
        except Exception:
            return

        for num in numbering.xpath("./*[local-name()='num']"):
            num_id = _xml_attr(num, "numId")
            abstract = _first_child(num, "abstractNumId")
            abstract_id = _xml_attr(abstract, "val") if abstract is not None else None
            if num_id and abstract_id:
                self._num_to_abstract[num_id] = abstract_id

        for abstract in numbering.xpath("./*[local-name()='abstractNum']"):
            abstract_id = _xml_attr(abstract, "abstractNumId")
            if not abstract_id:
                continue
            for lvl in abstract.xpath("./*[local-name()='lvl']"):
                ilvl = _xml_attr(lvl, "ilvl")
                if ilvl is None:
                    continue
                start_el = _first_child(lvl, "start")
                fmt_el = _first_child(lvl, "numFmt")
                text_el = _first_child(lvl, "lvlText")
                start = _safe_int(_xml_attr(start_el, "val"), 1)
                num_format = _xml_attr(fmt_el, "val") or "decimal"
                level_text = _xml_attr(text_el, "val") or ""
                self._levels[(abstract_id, ilvl)] = NumberingLevel(
                    start=start,
                    num_format=num_format,
                    level_text=level_text,
                )

    def _format_level_text(self, num_id: str, abstract_id: str, ilvl: str, level_text: str) -> str:
        label = level_text
        for placeholder in re.findall(r"%(\d+)", level_text):
            ref_ilvl = str(int(placeholder) - 1)
            ref_level = self._levels.get((abstract_id, ref_ilvl))
            if ref_level is None:
                continue
            value = self._counters.get((num_id, ref_ilvl), ref_level.start)
            label = label.replace(f"%{placeholder}", _format_number(value, ref_level.num_format))
        return re.sub(r"\s+", " ", label)


def _paragraph_numbering(paragraph) -> tuple[str | None, str | None]:
    num_pr = getattr(paragraph._p.pPr, "numPr", None) if paragraph._p.pPr is not None else None
    if num_pr is None:
        return None, None
    num_id = _xml_attr(num_pr.numId, "val") if num_pr.numId is not None else None
    ilvl = _xml_attr(num_pr.ilvl, "val") if num_pr.ilvl is not None else None
    return num_id, ilvl


def _xml_attr(element: Any, name: str) -> str | None:
    if element is None:
        return None
    return element.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{name}")


def _first_child(element: Any, local_name: str) -> Any | None:
    matches = element.xpath(f"./*[local-name()='{local_name}']")
    return matches[0] if matches else None


def _safe_int(value: str | None, default: int) -> int:
    try:
        return int(value) if value is not None else default
    except ValueError:
        return default


def _format_number(value: int, num_format: str) -> str:
    if num_format.lower().startswith("chinese"):
        return _to_chinese_number(value)
    return str(value)


def _to_chinese_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value <= 0:
        return str(value)
    if value < 10:
        return digits[value]
    if value == 10:
        return "十"
    if value < 20:
        return f"十{digits[value % 10]}"
    if value < 100:
        ten, one = divmod(value, 10)
        return f"{digits[ten]}十{digits[one] if one else ''}"
    return str(value)


def _find_body_start_index(paragraphs: list[DocxParagraph]) -> int:
    toc_index = next((i for i, paragraph in enumerate(paragraphs) if _is_toc_label(paragraph.text)), None)
    if toc_index is None:
        return 0

    headings: list[DocxHeading] = []
    for index, paragraph in enumerate(paragraphs[toc_index + 1 :], start=toc_index + 1):
        heading = _detect_heading(paragraph, index, use_native_headings=False)
        if heading:
            headings.append(heading)

    first_seen: dict[str, int] = {}
    for heading in headings:
        key = _heading_key(heading.title, heading.section_no)
        if key in first_seen:
            return heading.index
        first_seen.setdefault(key, heading.index)

    for index, paragraph in enumerate(paragraphs[toc_index + 1 :], start=toc_index + 1):
        if _is_toc_style(paragraph.style_id, paragraph.style_name) or _looks_like_toc_entry(paragraph.text):
            continue
        if _detect_heading(paragraph, index, use_native_headings=False):
            return index
    return toc_index + 1


def _detect_heading(paragraph: DocxParagraph, index: int, use_native_headings: bool) -> DocxHeading | None:
    native_level = _heading_level_from_style(paragraph.style_id, paragraph.style_name)
    if native_level in {1, 2, 3}:
        return DocxHeading(index=index, level=native_level, title=paragraph.text, section_no=_extract_section_no(paragraph.text))

    if use_native_headings:
        return None

    fallback = _numbered_heading(paragraph.text)
    if fallback is None:
        return None
    level, title, section_no = fallback
    return DocxHeading(index=index, level=level, title=title, section_no=section_no)


def _heading_level_from_style(style_id: str, style_name: str) -> int | None:
    combined = f"{style_id} {style_name}".lower()
    match = re.search(r"heading\s*(\d+)", combined)
    if match and int(match.group(1)) <= 3:
        return int(match.group(1))

    chinese_match = re.search(r"标题\s*(\d+)", f"{style_id} {style_name}")
    if chinese_match and int(chinese_match.group(1)) <= 3:
        return int(chinese_match.group(1))

    if "标题" in style_name and style_id in {"1", "2", "3"}:
        return int(style_id)
    return None


def _numbered_heading(text: str) -> tuple[int, str, str | None] | None:
    chapter_match = re.match(r"^\s*((?:第)?[一二三四五六七八九十百千万零〇两\d]+章)\s*(.+)?$", text)
    if chapter_match and _looks_like_title(text):
        return 1, text.strip(), chapter_match.group(1)

    number_match = re.match(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,2})\.?[、\s]+(.+)$", text)
    if not number_match or not _looks_like_title(text):
        return None

    section_no = number_match.group(1)
    level = section_no.count(".") + 1
    if level > 3:
        return None
    return level, text.strip(), section_no


def _extract_section_no(text: str) -> str | None:
    chapter_match = re.match(r"^\s*((?:第)?[一二三四五六七八九十百千万零〇两\d]+章)", text)
    if chapter_match:
        return chapter_match.group(1)

    number_match = re.match(r"^\s*(\d{1,2}(?:\.\d{1,2}){0,2})\.?", text)
    if number_match:
        return number_match.group(1)

    return None


def _is_standalone_chapter_no(text: str) -> bool:
    return re.match(r"^\s*(?:第)?[一二三四五六七八九十百千万零〇两\d]+章\s*$", text) is not None


def _merge_pending_chapter_heading(chapter: DocxHeading, heading: DocxHeading) -> DocxHeading:
    title = heading.title
    if not title.startswith(chapter.title):
        title = f"{chapter.title} {title}".strip()
    return DocxHeading(
        index=heading.index,
        level=chapter.level,
        title=title,
        section_no=chapter.section_no or chapter.title,
    )


def _is_toc_style(style_id: str, style_name: str) -> bool:
    combined = f"{style_id} {style_name}".lower()
    return "toc heading" in combined or re.search(r"\btoc\s*\d+\b", combined) is not None


def _is_toc_label(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text.strip())
    normalized = re.sub(r"^\d+[.．、]?", "", normalized)
    return normalized in {"目录", "目次", "contents"}


def _looks_like_toc_entry(text: str) -> bool:
    stripped = text.strip()
    return re.search(r"(\.{2,}|…+|·{2,})\s*[（(]?\d+[）)]?\s*$", stripped) is not None


def _heading_key(title: str, section_no: str | None) -> str:
    text = _strip_toc_page_suffix(title)
    if section_no:
        text = re.sub(rf"^\s*{re.escape(section_no)}\s*[.．、]?\s*", "", text)
        return f"{section_no}:{_normalize_text(text)}"
    return _normalize_text(text)


def _strip_toc_page_suffix(text: str) -> str:
    return re.sub(r"\s*(?:\.{2,}|…+|·{2,})\s*[（(]?\d+[）)]?\s*$", "", text.strip()).strip()


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", text).strip().lower()


def _looks_like_title(text: str) -> bool:
    stripped = text.strip()
    return bool(stripped) and len(stripped) <= 160 and not re.search(r"[。！？；;]$", stripped)
